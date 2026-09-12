"""Read publisher master data and monthly field service reports from the workbook."""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .models import MONTHS, PIONEER_STATUSES, MonthlyReport, Publisher

PUBINFO_SHEET = "PubInfo"
PUBINFO_FIRST_ROW = 2
MONTH_FIRST_ROW = 5

# Month sheets carry summary blocks below the publisher table; these labels mark the end of it.
_TRAILER_MARKERS = frozenset(
    {
        "name",
        "old reports",
        "totals",
        "publishers",
        "auxillary pioneers",
        "auxiliary pioneers",
        "regular pioneers",
        "special pioneers",
        "not reported",
        "grand total",
    }
)

_SERVICE_YEAR_RE = re.compile(r"(\d{4})\s*[-–]\s*(\d{4})")


class WorkbookError(RuntimeError):
    pass


def normalize_name(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def service_year_from_filename(path: Path) -> str:
    match = _SERVICE_YEAR_RE.search(path.stem)
    if not match:
        raise WorkbookError(
            f"Could not read a service year (e.g. 2026-2027) from the file name {path.name!r}. "
            "Pass --year explicitly."
        )
    return f"{match.group(1)}-{match.group(2)}"


def _text(value: object) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def _number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _is_trailer(name: str | None) -> bool:
    if name is None:
        return True
    lowered = name.casefold()
    return lowered in _TRAILER_MARKERS or lowered.startswith("old reports")


def read_publishers(workbook) -> tuple[list[Publisher], list[Publisher]]:
    """Return publishers to card, and those flagged 'Do not generate Card'."""
    if PUBINFO_SHEET not in workbook.sheetnames:
        raise WorkbookError(f"Workbook has no {PUBINFO_SHEET!r} sheet.")

    sheet = workbook[PUBINFO_SHEET]
    publishers: list[Publisher] = []
    excluded: list[Publisher] = []
    seen: set[str] = set()

    for row in sheet.iter_rows(min_row=PUBINFO_FIRST_ROW, max_col=9, values_only=True):
        name = _text(row[0])
        if not name:
            continue

        key = normalize_name(name)
        if key in seen:
            continue
        seen.add(key)

        pioneer_status = (_text(row[6]) or "").upper() or None
        if pioneer_status not in {"RP", "SP", "FM"}:
            pioneer_status = None

        publisher = Publisher(
            name=name,
            date_of_birth=_date(row[1]),
            date_of_baptism=_date(row[2]),
            gender=_text(row[3]),
            hope=_text(row[4]),
            appointment=_text(row[5]),
            pioneer_status=pioneer_status,
            field_service_group=_text(row[7]),
        )
        if _text(row[8]):
            excluded.append(publisher)
        else:
            publishers.append(publisher)

    return publishers, excluded


def read_month(workbook, month: str) -> dict[str, MonthlyReport]:
    """Read one month sheet. Rows below the first blank name (summary blocks) are ignored."""
    if month not in workbook.sheetnames:
        return {}

    sheet = workbook[month]
    reports: dict[str, MonthlyReport] = {}

    for row in sheet.iter_rows(min_row=MONTH_FIRST_ROW, max_col=9, values_only=True):
        name = _text(row[1])
        if _is_trailer(name):
            break

        status = (_text(row[3]) or "").upper() or None
        if status not in PIONEER_STATUSES:
            status = None

        reports[normalize_name(name)] = MonthlyReport(
            month=month,
            name=name,
            group=_text(row[2]),
            shared=(_text(row[4]) or "").upper() == "YES",
            bible_studies=int(_number(row[5]) or 0) if _number(row[5]) is not None else None,
            status=status,
            hours=_number(row[6]),
            remarks=_text(row[7]),
            credit=_number(row[8]),
        )

    return reports


def read_reports(workbook) -> dict[str, dict[str, MonthlyReport]]:
    """Map normalized publisher name -> month -> report."""
    by_publisher: dict[str, dict[str, MonthlyReport]] = {}
    for month in MONTHS:
        for key, report in read_month(workbook, month).items():
            by_publisher.setdefault(key, {})[month] = report
    return by_publisher


def load(path: Path):
    workbook = load_workbook(path, data_only=True, read_only=True, keep_vba=False)
    try:
        publishers, excluded = read_publishers(workbook)
        reports = read_reports(workbook)
    finally:
        workbook.close()
    return publishers, excluded, reports
